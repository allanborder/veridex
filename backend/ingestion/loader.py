import os
import csv
from pathlib import Path
from typing import List, Tuple
from pypdf import PdfReader

try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

class DocumentLoader:
    """
    Safely loads documents (.txt, .md, .pdf, .docx, .xlsx, .csv) from local directory.
    Zero external network telemetry.
    """
    SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".xlsx", ".csv"}

    @classmethod
    def load_file(cls, file_path: Path) -> Tuple[str, str]:
        """Returns (cleaned_text, filename)"""
        if not file_path.exists():
            raise FileNotFoundError(f"Document not found: {file_path}")

        ext = file_path.suffix.lower()
        if ext not in cls.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: {ext}")

        filename = file_path.name

        if ext in {".txt", ".md"}:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        elif ext == ".pdf":
            reader = PdfReader(str(file_path))
            pages_text = []
            for i, page in enumerate(reader.pages):
                extracted = page.extract_text()
                if extracted:
                    pages_text.append(f"Page {i+1}:\n{extracted}")
            text = "\n\n".join(pages_text)
        elif ext == ".docx":
            if not docx:
                raise ValueError("DOCX parser not installed.")
            doc = docx.Document(str(file_path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n\n".join(paragraphs)
        elif ext == ".xlsx":
            if not openpyxl:
                raise ValueError("Excel parser not installed.")
            wb = openpyxl.load_workbook(str(file_path), data_only=True)
            sheet_texts = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                rows = []
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(v) for v in row if v is not None]
                    if row_vals:
                        rows.append(" | ".join(row_vals))
                if rows:
                    sheet_texts.append(f"Sheet: {sheet}\n" + "\n".join(rows))
            text = "\n\n".join(sheet_texts)
        elif ext == ".csv":
            rows = []
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                for row in reader:
                    if row:
                        rows.append(" | ".join(row))
            text = "\n".join(rows)
        else:
            text = ""

        return text.strip(), filename

    @classmethod
    def load_directory(cls, dir_path: Path) -> List[Tuple[str, str, Path]]:
        """Loads all supported documents in directory"""
        if not dir_path.exists():
            return []

        results = []
        for p in dir_path.glob("*"):
            if p.is_file() and p.suffix.lower() in cls.SUPPORTED_EXTENSIONS:
                try:
                    text, fname = cls.load_file(p)
                    if text:
                        results.append((text, fname, p))
                except Exception as e:
                    print(f"[Loader] Error reading {p.name}: {e}")
        return results
